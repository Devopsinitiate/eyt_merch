from django.contrib import admin
from django.contrib.auth.admin import UserAdmin
from django.http import HttpResponse
from django.contrib import messages
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from .models import Size, Order, CustomUser, TournamentConfig, TournamentRegistration

@admin.action(description='Confirm payment for selected registrations')
def confirm_payment(modeladmin, request, queryset):
    """Mark selected registrations as payment confirmed"""
    updated = queryset.update(payment_status='CONFIRMED')
    messages.success(request, f'{updated} registration(s) marked as payment confirmed.')


@admin.action(description='Export selected registrations to Excel')
def export_tournament_registrations(modeladmin, request, queryset):
    """Export tournament registrations to Excel file"""
    wb = Workbook()
    ws = wb.active
    ws.title = "Tournament Registrations"

    # Header styling
    header_fill = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True, size=12)

    # Headers
    headers = ['Reg ID', 'Full Name', 'Gamer Tag', 'Gender', 'Email', 'Payment Reference', 'Payment Status', 'Date']
    ws.append(headers)

    # Style header row
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center')

    # Data rows
    for reg in queryset:
        ws.append([
            reg.id,
            reg.full_name,
            reg.gamer_tag,
            reg.get_gender_display(),
            reg.email,
            reg.payment_reference,
            reg.get_payment_status_display(),
            reg.created_at.strftime('%Y-%m-%d %H:%M')
        ])

    # Adjust column widths
    column_widths = [10, 25, 20, 18, 30, 25, 15, 20]
    for i, width in enumerate(column_widths, 1):
        ws.column_dimensions[chr(64 + i)].width = width

    # Create response
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = 'attachment; filename=eyt_tournament_registrations.xlsx'
    wb.save(response)
    return response


@admin.action(description='Export selected orders to Excel')
def export_to_excel(modeladmin, request, queryset):
    """Export orders to Excel file"""
    wb = Workbook()
    ws = wb.active
    ws.title = "Orders"
    
    # Header styling
    header_fill = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True, size=12)
    
    # Headers
    headers = ['Order ID', 'User Email', 'Gamer Tag', 'Name', 'Number', 'Size', 'Color', 'Phone', 'Date']
    ws.append(headers)
    
    # Style header row
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center')
    
    # Data rows
    for order in queryset:
        ws.append([
            order.id,
            order.user.email if order.user else order.email,
            order.user.gamer_tag if order.user else order.gamer_tag,
            order.user.full_name if order.user else order.full_name,
            order.preferred_number,
            order.size,
            order.get_color_display(),
            order.user.phone if order.user else order.phone,
            order.created_at.strftime('%Y-%m-%d %H:%M')
        ])
    
    # Adjust column widths
    column_widths = [10, 30, 20, 25, 10, 10, 20, 15, 20]
    for i, width in enumerate(column_widths, 1):
        ws.column_dimensions[chr(64 + i)].width = width
    
    # Create response
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = 'attachment; filename=eyt_orders.xlsx'
    wb.save(response)
    return response


@admin.register(CustomUser)
class CustomUserAdmin(UserAdmin):
    list_display = ['email', 'gamer_tag', 'full_name', 'is_active', 'date_joined']
    list_filter = ['is_active', 'is_staff', 'date_joined']
    search_fields = ['email', 'gamer_tag', 'full_name', 'phone']
    ordering = ['-date_joined']
    
    fieldsets = (
        (None, {'fields': ('email', 'password')}),
        ('EYT Information', {'fields': ('gamer_tag', 'full_name', 'phone')}),
        ('Permissions', {
            'fields': ('is_active', 'is_staff', 'is_superuser', 'groups', 'user_permissions'),
            'classes': ('collapse',)
        }),
        ('Important dates', {'fields': ('last_login', 'date_joined')}),
    )
    
    add_fieldsets = (
        (None, {
            'classes': ('wide',),
            'fields': ('email', 'gamer_tag', 'full_name', 'phone', 'password1', 'password2'),
        }),
    )


@admin.register(Size)
class SizeAdmin(admin.ModelAdmin):
    list_display = ['name', 'available_quantity', 'is_available']
    list_editable = ['available_quantity']
    ordering = ['name']


@admin.register(Order)
class OrderAdmin(admin.ModelAdmin):
    list_display = ['id', 'user_gamer_tag', 'full_name', 'size', 'color', 'created_at']
    list_filter = ['size', 'color', 'created_at']
    search_fields = ['full_name', 'gamer_tag', 'user__email', 'user__gamer_tag']
    readonly_fields = ['created_at', 'updated_at']
    actions = [export_to_excel]
    
    def user_gamer_tag(self, obj):
        return obj.user.gamer_tag if obj.user else obj.gamer_tag
    user_gamer_tag.short_description = 'Gamer Tag'
    
    fieldsets = (
        ('User', {
            'fields': ('user',)
        }),
        ('Personal Information', {
            'fields': ('full_name', 'gamer_tag', 'preferred_number')
        }),
        ('Product Selection', {
            'fields': ('size', 'color')
        }),
        ('Contact Information', {
            'fields': ('phone', 'email')
        }),
        ('Metadata', {
            'fields': ('created_at', 'updated_at'),
            'classes': ('collapse',)
        }),
    )


@admin.register(TournamentConfig)
class TournamentConfigAdmin(admin.ModelAdmin):
    list_display = [
        'tournament_name', 'fee_amount', 'currency', 'account_number',
        'bank_name', 'spots', 'is_registration_open'
    ]
    ordering = ['tournament_name']

    def has_add_permission(self, request):
        # A single config row is pre-seeded; prevent accidental duplicates.
        return False

    def spots(self, obj):
        if obj.spots_remaining is None:
            return f"{obj.registered_count} / Unlimited"
        return f"{obj.registered_count} / {obj.max_participants}"
    spots.short_description = 'Spots'


@admin.register(TournamentRegistration)
class TournamentRegistrationAdmin(admin.ModelAdmin):
    list_display = [
        'id', 'gamer_tag', 'full_name', 'gender', 'email',
        'get_payment_status_display', 'created_at'
    ]
    list_filter = ['gender', 'payment_status', 'created_at']
    search_fields = ['full_name', 'gamer_tag', 'email', 'payment_reference']
    readonly_fields = ['created_at']
    actions = [confirm_payment, export_tournament_registrations]

    def get_payment_status_display(self, obj):
        return obj.get_payment_status_display()
    get_payment_status_display.short_description = 'Payment Status'
    get_payment_status_display.admin_order_field = 'payment_status'
